"""Export key analytics outputs to Excel workbook."""
from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

from ..config import ANALYTICS_DIR, EXCEL_DIR, REPORT_DIR

logger = logging.getLogger(__name__)


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")


def export_to_excel() -> Path:
    out = EXCEL_DIR / "banking_analytics.xlsx"
    wb = Workbook()
    wb.remove(wb.active)

    sheets = [
        ("Customers", ANALYTICS_DIR / "2_curated_customers.csv"),
        ("Accounts", ANALYTICS_DIR / "1_curated_accounts.csv"),
        ("Branch Performance", ANALYTICS_DIR / "3_branch_performance.csv"),
        ("Loan Book", ANALYTICS_DIR / "7_loan_book.csv"),
        ("Loan Payments", ANALYTICS_DIR / "8_loan_payments.csv"),
        ("Support Tickets", ANALYTICS_DIR / "9_support_ticket_summary.csv"),
        ("RFM Segments", ANALYTICS_DIR / "10_customer_rfm_segments.csv"),
        ("Credit Risk", ANALYTICS_DIR / "11_loan_default_features.csv"),
        ("Fraud Summary", ANALYTICS_DIR / "12_fraud_card_summary.csv"),
        ("Credit Calibration", ANALYTICS_DIR / "credit_calibration.csv"),
        ("Credit Cutoffs", ANALYTICS_DIR / "credit_cutoffs.csv"),
    ]

    for sheet_name, csv_path in sheets:
        if not csv_path.exists():
            logger.warning("SKIP missing %s", csv_path.name)
            continue
        df = pd.read_csv(csv_path)
        ws = wb.create_sheet(title=sheet_name[:31])
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        _style_header(ws)
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    wb.save(out)
    logger.info("Excel -> %s (%s sheets)", out.name, len(wb.sheetnames))
    return out


if __name__ == "__main__":
    export_to_excel()
